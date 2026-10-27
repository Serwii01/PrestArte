"""Genera la presentación comercial de PrestArte (16:9, editable).

Uso: python3 presentacion.py  ->  ../PrestArte_Presentacion.pptx
Para el PDF: soffice --headless --convert-to pdf ../PrestArte_Presentacion.pptx --outdir ..
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Mm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

AQUI = Path(__file__).parent
CAP = AQUI.parent / "capturas"
LOGO = AQUI / "logo.png"
SALIDA = AQUI.parent / "PrestArte_Presentacion.pptx"

AZUL = RGBColor(0x11, 0x52, 0xD4)
AZUL_OSC = RGBColor(0x0B, 0x2F, 0x86)
AZUL_CLARO = RGBColor(0xDB, 0xE1, 0xFF)
FONDO_SUAVE = RGBColor(0xF3, 0xF3, 0xFE)
TINTA = RGBColor(0x19, 0x1B, 0x23)
GRIS = RGBColor(0x43, 0x46, 0x54)
GRIS_CLARO = RGBColor(0x73, 0x76, 0x86)
LINEA = RGBColor(0xE2, 0xE1, 0xED)
TERRA = RGBColor(0xA6, 0x38, 0x00)
TERRA_CLARO = RGBColor(0xFF, 0xF1, 0xEB)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
FUENTE = "Inter"

prs = Presentation()
prs.slide_width, prs.slide_height = Mm(338.67), Mm(190.5)
W, H = prs.slide_width, prs.slide_height
VACIA = prs.slide_layouts[6]
M = Mm(18)  # margen lateral


def texto(slide, x, y, w, h, parrafos, tam=14, color=TINTA, negrita=False, alinear=PP_ALIGN.LEFT,
          ancla=MSO_ANCHOR.TOP, interlinea=1.15):
    """Caja de texto. `parrafos` es un str o una lista de str / (str, dict de estilo)."""
    caja = slide.shapes.add_textbox(x, y, w, h)
    tf = caja.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ancla
    if isinstance(parrafos, str):
        parrafos = [parrafos]
    for i, par in enumerate(parrafos):
        estilo = {}
        if isinstance(par, tuple):
            par, estilo = par
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinear
        p.line_spacing = interlinea
        p.space_after = Pt(estilo.get("despues", 6))
        r = p.add_run()
        r.text = par
        f = r.font
        f.name = FUENTE
        t = estilo.get("tam", tam)
        f.size = Pt(t * 1.2 if 10 < t <= 15 else t)  # cuerpo algo mayor para proyectar
        f.bold = estilo.get("negrita", negrita)
        f.color.rgb = estilo.get("color", color)
    return caja


def rect(slide, x, y, w, h, relleno=FONDO_SUAVE, borde=None, redondeo=0.06):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.adjustments[0] = redondeo
    s.fill.solid()
    s.fill.fore_color.rgb = relleno
    if borde:
        s.line.color.rgb = borde
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    sp = s._element.spPr
    from pptx.oxml.ns import qn
    if sp.find(qn("a:effectLst")) is None:
        sp.append(sp.makeelement(qn("a:effectLst"), {}))
    return s


def imagen(slide, ruta, x, y, w=None, h=None):
    pic = slide.shapes.add_picture(str(ruta), x, y, w, h)
    pic.line.color.rgb = LINEA
    pic.line.width = Pt(0.75)
    return pic


def cabecera(slide, antetitulo, titulo, n):
    texto(slide, M, Mm(14), W - 2 * M, Mm(6), antetitulo.upper(), tam=10, color=AZUL, negrita=True)
    texto(slide, M, Mm(21), W - 2 * M, Mm(16), titulo, tam=27, negrita=True)
    texto(slide, M, H - Mm(10), Mm(120), Mm(5), "PrestArte · Presentación confidencial", tam=8, color=GRIS_CLARO)
    texto(slide, W - M - Mm(20), H - Mm(10), Mm(20), Mm(5), str(n), tam=8, color=GRIS_CLARO,
          alinear=PP_ALIGN.RIGHT)


def notas(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


def tarjeta(slide, x, y, w, h, titulo, cuerpo, relleno=FONDO_SUAVE, tam=11.5, borde=None, color_titulo=TINTA):
    rect(slide, x, y, w, h, relleno, borde)
    pad = Mm(5)
    items = [(titulo, {"tam": 13.5, "negrita": True, "despues": 6, "color": color_titulo})]
    if isinstance(cuerpo, str):
        cuerpo = [cuerpo]
    items += [(c, {"tam": tam, "color": GRIS, "despues": 4}) for c in cuerpo]
    texto(slide, x + pad, y + pad, w - 2 * pad, h - 2 * pad, items)


def cifra(slide, x, y, w, valor, etiqueta, color=AZUL, relleno=FONDO_SUAVE, h=Mm(38)):
    rect(slide, x, y, w, h, relleno)
    texto(slide, x + Mm(5), y + Mm(5), w - Mm(10), Mm(14), valor, tam=30, negrita=True, color=color)
    texto(slide, x + Mm(5), y + Mm(20), w - Mm(10), h - Mm(22), etiqueta, tam=11, color=GRIS)


def columnas(n, gap=Mm(6), x0=M, ancho=None):
    ancho = ancho or (W - 2 * M)
    w = (ancho - gap * (n - 1)) / n
    return [int(x0 + i * (w + gap)) for i in range(n)], int(w)


n = 0


def nueva():
    global n
    n += 1
    return prs.slides.add_slide(VACIA), n


# ---------- 1. Portada ----------
s, _ = nueva()
fondo = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
fondo.fill.solid(); fondo.fill.fore_color.rgb = AZUL_OSC; fondo.line.fill.background()
banda = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(W * 0.58), 0, int(W * 0.42), H)
banda.fill.solid(); banda.fill.fore_color.rgb = AZUL; banda.line.fill.background()
rect(s, M, Mm(18), Mm(62), Mm(20), BLANCO, redondeo=0.12)
s.shapes.add_picture(str(LOGO), M + Mm(6), Mm(22), height=Mm(12))
texto(s, M, Mm(62), Mm(170), Mm(60), [
    ("La plataforma para prestar arte privado a museos y fundaciones", {"tam": 34, "negrita": True, "color": BLANCO, "despues": 14}),
    ("Coleccionistas, instituciones y transportistas especializados gestionan cada préstamo en un solo lugar: "
     "solicitud, seguro, permisos, transporte, exposición y devolución.", {"tam": 15, "color": AZUL_CLARO}),
])
texto(s, M, H - Mm(24), Mm(170), Mm(10), "Oportunidad de adquisición · Octubre de 2026 · Confidencial", tam=11, color=AZUL_CLARO)
imagen(s, CAP / "06_prestamo_transito.jpg", int(W * 0.60), Mm(40), w=int(W * 0.37))
notas(s, "Presentación de PrestArte para posibles compradores. Mensaje principal: es un producto completo y "
         "funcional que resuelve la gestión de préstamos de arte privado, con el conocimiento del sector ya "
         "convertido en software.")

# ---------- 2. Problema ----------
s, k = nueva()
cabecera(s, "El problema", "Prestar una obra privada hoy es lento, opaco y arriesgado", k)
xs, w = columnas(3)
y, h = Mm(48), Mm(108)
tarjeta(s, xs[0], y, w, h, "El coleccionista", [
    "No sabe qué instituciones podrían interesarse por su obra.",
    "Teme exponer su identidad, el valor y la ubicación.",
    "No controla quién transporta la obra ni cómo.",
    "Recibe poco a cambio: ni mención ni historial."])
tarjeta(s, xs[1], y, w, h, "El museo o la fundación", [
    "El inventario privado disponible es invisible.",
    "Los costes de transporte, caja, seguro y correo llegan tarde.",
    "Seguros, permisos de exportación y procedencia se gestionan a mano.",
    "No hay un expediente único y auditable."])
tarjeta(s, xs[2], y, w, h, "El transportista", [
    "Los encargos llegan por correo y con datos incompletos.",
    "Presupuestar y coordinar recogidas exige mucho trabajo manual.",
    "Documentar el estado de la obra en cada hito es engorroso."])
notas(s, "Hoy un préstamo de una colección privada se organiza con correos, llamadas y hojas de cálculo. "
         "Cada una de las tres partes sufre un problema distinto, y nadie tiene la foto completa.")

# ---------- 3. Oportunidad ----------
s, k = nueva()
cabecera(s, "La oportunidad", "Un mercado grande con un proceso que sigue siendo artesanal", k)
xs, w = columnas(3)
cifra(s, xs[0], Mm(50), w, "59.600 M$", "Ventas del mercado mundial del arte en 2025, un 4 % más que en 2024 (Art Basel y UBS)")
cifra(s, xs[1], Mm(50), w, "~1.500", "Museos y colecciones museográficas en España (Ministerio de Cultura)")
cifra(s, xs[2], Mm(50), w, "+65 M", "Visitas anuales a esas instituciones (estadística oficial de 2018)")
texto(s, M, Mm(100), W - 2 * M, Mm(40), [
    "Las exposiciones temporales dependen de los préstamos, y una parte importante del patrimonio está en manos privadas. "
    "Las herramientas habituales de gestión de colecciones trabajan dentro de cada institución; PrestArte es el punto "
    "de encuentro entre las tres partes del préstamo, con el cumplimiento normativo dentro del flujo."], tam=15, color=GRIS)
texto(s, M, Mm(160), W - 2 * M, Mm(10), "Fuentes: Art Basel y UBS, The Art Market (2025 y 2026); Ministerio de Cultura, "
      "Estadística de Museos y Colecciones Museográficas.", tam=9, color=GRIS_CLARO)
notas(s, "Las cifras son de contexto, no un tamaño de mercado direccionable. Si el comprador pide un cálculo de "
         "mercado, conviene hacerlo con sus propios datos de clientes (transportistas, aseguradoras o museos).")

# ---------- 4. Solución ----------
s, k = nueva()
cabecera(s, "La solución", "Una red cerrada y verificada, con un flujo para cada actor", k)
xs, w = columnas(4, gap=Mm(5))
perfiles = [
    ("Coleccionista", "Publica sus obras con requisitos de conservación y condiciones. Decide a quién presta y elige transportista."),
    ("Museo / fundación", "Crea exposiciones, busca obras, ve costes y CO₂ antes de pedir y aprueba presupuestos."),
    ("Transportista", "Perfil público con cobertura y tarifas. Cotiza ida y vuelta y registra cada hito con fotos."),
    ("Administración", "Verifica cuentas (KYB), revisa la procedencia de las obras y supervisa la actividad."),
]
for x, (t, c) in zip(xs, perfiles):
    tarjeta(s, x, Mm(48), w, Mm(74), t, c, tam=12)
rect(s, M, Mm(130), W - 2 * M, Mm(30), AZUL_CLARO)
texto(s, M + Mm(6), Mm(134), W - 2 * M - Mm(12), Mm(24),
      "Nadie opera sin que la administración haya revisado antes su documentación (DNI, NIE o CIF y un documento "
      "acreditativo). Cada préstamo sigue una máquina de estados: no se puede saltar un paso, y cada acción solo la "
      "puede hacer la parte a la que le toca.", tam=13.5, color=AZUL_OSC, ancla=MSO_ANCHOR.MIDDLE)
notas(s, "La red cerrada y verificada es la base de la confianza: sin ella, ningún coleccionista publicaría una obra valiosa.")

# ---------- 5. Flujo ----------
s, k = nueva()
cabecera(s, "Cómo funciona", "Un préstamo, de la solicitud a la devolución", k)
pasos = [
    ("Solicitud", "La institución pide la obra para una exposición, con fechas y sede."),
    ("Aceptación", "El coleccionista acepta, fija el valor asegurado y elige transportista."),
    ("Presupuesto", "El transportista cotiza ida y vuelta; la institución lo aprueba."),
    ("Seguro", "Seguro clavo a clavo y, si procede, permiso de exportación."),
    ("Envío", "Recogida, tránsito y entrega con informe de estado y fotos."),
    ("Exposición y vuelta", "Ampliación opcional, devolución y valoraciones."),
]
xs, w = columnas(6, gap=Mm(4))
for i, (x, (t, c)) in enumerate(zip(xs, pasos)):
    rect(s, x, Mm(50), w, Mm(78), FONDO_SUAVE)
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Mm(5), Mm(56), Mm(12), Mm(12))
    circ.fill.solid(); circ.fill.fore_color.rgb = AZUL; circ.line.fill.background()
    circ.text_frame.text = str(i + 1)
    para = circ.text_frame.paragraphs[0]; para.alignment = PP_ALIGN.CENTER
    para.runs[0].font.size = Pt(14); para.runs[0].font.bold = True; para.runs[0].font.name = FUENTE
    para.runs[0].font.color.rgb = BLANCO
    texto(s, x + Mm(5), Mm(73), w - Mm(10), Mm(52), [(t, {"tam": 13.5, "negrita": True}), (c, {"tam": 11, "color": GRIS})])
texto(s, M, Mm(138), W - 2 * M, Mm(30), [
    "Se genera automáticamente: contrato en PDF, certificado de exposición, historial de eventos con autor y fecha, "
    "avisos en la aplicación y por email, y recordatorios diarios de plazos (fin del préstamo, seguro, permiso)."],
    tam=13.5, color=GRIS)
notas(s, "La obra no puede salir sin el seguro aprobado y, si sale de España y está protegida, sin permiso de "
         "exportación. Un presupuesto rechazado permite cambiar de transportista sin perder el histórico.")

# ---------- 6. Catálogo ----------
s, k = nueva()
cabecera(s, "El producto", "Un catálogo pensado para instituciones", k)
pic = imagen(s, CAP / "02_catalogo.jpg", M, Mm(44), h=Mm(126))
x2 = pic.left + pic.width + Mm(8)
texto(s, x2, Mm(48), W - x2 - M, Mm(120), [
    ("Búsqueda con filtros", {"tam": 14, "negrita": True}),
    ("Texto, ciudad, años, estado, procedencia verificada y obras libres entre dos fechas.", {"tam": 12, "color": GRIS, "despues": 12}),
    ("Calendario de disponibilidad", {"tam": 14, "negrita": True}),
    ("Días ocupados visibles en la ficha y en el calendario de la solicitud.", {"tam": 12, "color": GRIS, "despues": 12}),
    ("Visibilidad controlada", {"tam": 14, "negrita": True}),
    ("Cada obra puede ser pública o solo para instituciones verificadas.", {"tam": 12, "color": GRIS, "despues": 12}),
    ("Capturas reales de la aplicación con datos de demostración.", {"tam": 9.5, "color": GRIS_CLARO}),
])
notas(s, "Todas las capturas son de la aplicación en funcionamiento (octubre de 2026), no de prototipos.")

# ---------- 7. Ficha + solicitud ----------
s, k = nueva()
cabecera(s, "El producto", "Antes de pedir, la institución sabe si cumple y cuánto costará", k)
xs, w = columnas(2)
imagen(s, CAP / "03_ficha_obra.jpg", xs[0], Mm(44), w=w)
imagen(s, CAP / "10_solicitud.jpg", xs[1], Mm(44), w=w)
texto(s, xs[0], Mm(150), w, Mm(20), "Ficha de la obra: coleccionista anónimo y valor visible solo para instituciones verificadas.", tam=11, color=GRIS)
texto(s, xs[1], Mm(150), w, Mm(20), "Solicitud: requisitos de la obra frente a las instalaciones del museo, con aviso de cada incumplimiento.", tam=11, color=GRIS)
notas(s, "La compatibilidad no bloquea (decide el coleccionista), salvo la antelación mínima. También se muestra "
         "una estimación de coste y huella de carbono con la tarifa del transportista o una de referencia.")

# ---------- 8. Paneles ----------
s, k = nueva()
cabecera(s, "El producto", "Cada perfil tiene su panel con lo que le toca hacer", k)
xs, w = columnas(2)
imagen(s, CAP / "05_panel_coleccionista.jpg", xs[0], Mm(44), w=w)
imagen(s, CAP / "13_panel_transporte.jpg", xs[1], Mm(44), w=w)
texto(s, xs[0], Mm(150), w, Mm(20), "Coleccionista: obras, préstamos activos, solicitudes pendientes y valor asegurado.", tam=11, color=GRIS)
texto(s, xs[1], Mm(150), w, Mm(20), "Transportista: envíos con número de seguimiento, estado y precio.", tam=11, color=GRIS)
notas(s, "También existen el panel de la fundación (solicitudes, exposiciones, presupuestos por aprobar) y el de "
         "administración (altas pendientes y diligencia debida).")

# ---------- 9. Expediente ----------
s, k = nueva()
cabecera(s, "El producto", "Un expediente único y trazable para cada préstamo", k)
pic = imagen(s, CAP / "06_prestamo_transito.jpg", M, Mm(44), h=Mm(126))
x2 = pic.left + pic.width + Mm(8)
texto(s, x2, Mm(48), W - x2 - M, Mm(120), [
    ("Línea de estados", {"tam": 14, "negrita": True}),
    ("14 estados con transiciones controladas.", {"tam": 12, "color": GRIS, "despues": 10}),
    ("Informes de estado", {"tam": 14, "negrita": True}),
    ("Fotos obligatorias en cada recogida y entrega.", {"tam": 12, "color": GRIS, "despues": 10}),
    ("Documentos", {"tam": 14, "negrita": True}),
    ("Contrato, seguro, permiso y certificado de exposición.", {"tam": 12, "color": GRIS, "despues": 10}),
    ("Chat por préstamo", {"tam": 14, "negrita": True}),
    ("Entre las tres partes, con imágenes y PDF.", {"tam": 12, "color": GRIS}),
])
notas(s, "La dirección de recogida es privada: solo la ve el transportista cuando tiene un envío aprobado y en curso.")

# ---------- 10. Diferenciales ----------
s, k = nueva()
cabecera(s, "Ventajas competitivas", "El conocimiento del sector ya convertido en software", k)
difs = [
    ("Cumplimiento normativo", "Seguro clavo a clavo y permiso de exportación (Ley 16/1985). Sin ellos, la obra no sale."),
    ("Procedencia", "Historial de propietarios, aviso de huecos en 1933–1945 y diligencia debida contra bases de obras robadas."),
    ("Privacidad", "Coleccionista anónimo, valor en horquilla y dirección de recogida protegida."),
    ("Costes y CO₂", "Estimación previa de transporte, caja, seguro y correo, y de la huella de carbono."),
    ("Incentivos", "Mención en cartela, usos de imagen, canon opcional y certificado de exposición."),
    ("Trazabilidad", "Informes de estado con fotos, correo virtual e historial de cada evento."),
]
xs, w = columnas(3)
for i, (t, c) in enumerate(difs):
    tarjeta(s, xs[i % 3], Mm(46) + (i // 3) * Mm(60), w, Mm(54), t, c, tam=12)
notas(s, "Un marketplace genérico se hace en semanas; esto es lo que cuesta copiar, porque requiere conocer el "
         "préstamo de arte: seguros, normativa de patrimonio y procedencia.")

# ---------- 11. Modelo de negocio ----------
s, k = nueva()
cabecera(s, "Modelo de negocio", "Varias vías de ingresos sobre la misma plataforma", k)
vias = [
    ("Comisión por préstamo", "Sobre cada préstamo cerrado. El paso «Pagado» ya existe en el flujo."),
    ("Suscripción institucional", "Cuota anual para museos y fundaciones."),
    ("Plan para transportistas", "Directorio, tarifas publicadas y encargos cualificados."),
    ("Verificación de procedencia", "Servicio de diligencia debida con distintivo."),
    ("Seguro intermediado", "Con un socio asegurador con licencia de mediación."),
    ("Marca blanca", "Instancia propia para una red de museos o un transportista."),
]
xs, w = columnas(3)
for i, (t, c) in enumerate(vias):
    tarjeta(s, xs[i % 3], Mm(46) + (i // 3) * Mm(46), w, Mm(40), t, c, tam=12)
texto(s, M, Mm(144), W - 2 * M, Mm(20),
      "PrestArte aún no factura ni tiene pasarela de pago: son propuestas que la arquitectura actual admite sin rediseño.",
      tam=12, color=GRIS_CLARO)
notas(s, "Dejar claro que no hay ingresos. Preguntar al comprador qué vía encaja con su negocio: un transportista "
         "querrá el plan profesional o la marca blanca; una aseguradora, el seguro intermediado.")

# ---------- 12. Tecnología ----------
s, k = nueva()
cabecera(s, "Tecnología y estado", "Base técnica estándar, verificada en octubre de 2026", k)
xs, w = columnas(4, gap=Mm(5))
for x, (v, e) in zip(xs, [("102", "endpoints REST con permisos por operación"), ("31", "entidades de dominio"),
                         ("~27.500", "líneas de código propio"), ("0 €", "en licencias: todo es software libre")]):
    cifra(s, x, Mm(46), w, v, e, h=Mm(36))
xs2, w2 = columnas(2)
tarjeta(s, xs2[0], Mm(90), w2, Mm(72), "Stack", [
    "Frontend: Angular 21 y Tailwind CSS 4",
    "Backend: Java 21, Spring Boot 4, Spring Security, JPA",
    "Base de datos: MySQL 8",
    "Seguridad: JWT, BCrypt, límite de intentos, archivos protegidos",
    "Documentos: OpenPDF y correo SMTP"], tam=12, relleno=BLANCO, borde=LINEA)
tarjeta(s, xs2[1], Mm(90), w2, Mm(72), "Verificado", [
    "El backend compila con Java 21 y arranca contra MySQL/MariaDB.",
    "Préstamo completo por API, de la solicitud a la devolución: todos los pasos correctos.",
    "Frontend recorrido con los cuatro perfiles.",
    "Ficha técnica con arquitectura, API y deuda técnica."], tam=12)
notas(s, "Si el comprador tiene equipo técnico, entregarle la ficha técnica y ofrecer una demo de instalación. "
         "La deuda técnica está listada con honestidad: tests, migraciones, Docker y textos legales.")

# ---------- 13. Hoja de ruta ----------
s, k = nueva()
cabecera(s, "Hoja de ruta", "Lo que falta es poner en producción, no construir", k)
xs, w = columnas(3)
tarjeta(s, xs[0], Mm(46), w, Mm(112), "Hecho", [
    "Ciclo completo del préstamo y del envío", "Verificación de cuentas", "Seguro, permisos y procedencia",
    "Costes, CO₂ y compatibilidad", "Chat, avisos y recordatorios", "Contratos y certificados PDF"], tam=12)
tarjeta(s, xs[1], Mm(46), w, Mm(112), "Antes de lanzar", [
    "Tests automáticos", "Migraciones (Flyway) y Docker", "Despliegue y monitorización",
    "Archivos en almacenamiento de objetos", "Textos legales (RGPD, cookies)", "Calibrar tarifas de referencia"],
    tam=12, relleno=TERRA_CLARO, color_titulo=TERRA)
tarjeta(s, xs[2], Mm(46), w, Mm(112), "Crecimiento", [
    "Pasarela de pago y comisión", "Firma electrónica (eIDAS)", "Inglés y otros idiomas",
    "Avisos en tiempo real", "Integración con aseguradoras", "App móvil para el transportista"], tam=12)
notas(s, "Riesgos a reconocer: no hay usuarios ni ingresos, y un marketplace necesita un actor ancla que aporte el "
         "primer lado del mercado. Por eso encaja mejor con un comprador que ya esté en el sector.")

# ---------- 14. Operación ----------
s, k = nueva()
cabecera(s, "La operación", "Qué incluye y cómo seguimos", k)
xs, w = columnas(2)
tarjeta(s, xs[0], Mm(46), w, Mm(84), "Incluye", [
    "Código fuente completo (frontend y backend) con su historial",
    "Marca PrestArte, logotipo y diseños",
    "Documentación: memoria técnica, README, ficha técnica y dosier",
    "Datos de demostración",
    "Transición y soporte durante [X semanas]"], tam=12.5, relleno=BLANCO, borde=LINEA)
tarjeta(s, xs[1], Mm(46), w, Mm(84), "Próximos pasos", [
    "1. Acuerdo de confidencialidad", "2. Demostración en directo", "3. Acceso al código y a la ficha técnica",
    "4. Revisión técnica y legal", "5. Oferta, contrato y traspaso"], tam=12.5)
rect(s, M, Mm(138), W - 2 * M, Mm(26), AZUL_OSC)
texto(s, M + Mm(8), Mm(141), W - 2 * M - Mm(16), Mm(20), [
    ("Sergio Fernández Morales", {"tam": 16, "negrita": True, "color": BLANCO, "despues": 2}),
    ("[email] · [teléfono] · Abierto a venta total o licencia", {"tam": 13, "color": AZUL_CLARO}),
], ancla=MSO_ANCHOR.MIDDLE)
notas(s, "Cerrar proponiendo la demo en directo y el acuerdo de confidencialidad. Sustituir los textos entre "
         "corchetes antes de enviar.")

prs.save(SALIDA)
print("Generado", SALIDA)
